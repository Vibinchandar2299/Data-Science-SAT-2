"""
generate_presentation.py
Generates the executive PowerPoint presentation: outputs/business_insights_pitch.pptx
With 16:9 widescreen layout, modern corporate aesthetics, KPI cards, and embedded charts.
Also exports the presentation to PDF via Windows PowerPoint COM automation.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

OUTPUT_PPTX = os.path.join("outputs", "business_insights_pitch.pptx")
OUTPUT_PDF = os.path.join("outputs", "business_insights_pitch.pdf")
FIG_DIR = os.path.join("outputs", "figures")
os.makedirs("outputs", exist_ok=True)

prs = Presentation()
# Set 16:9 widescreen dimensions
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_slide_layout = prs.slide_layouts[6]

# Colors
COLOR_NAVY = RGBColor(27, 54, 93)      # #1B365D - Primary Header
COLOR_DARK = RGBColor(17, 24, 39)      # #111827 - Body Text
COLOR_MUTED = RGBColor(100, 116, 139)  # #64748B - Subtitles/Notes
COLOR_CORAL = RGBColor(231, 76, 60)    # #E74C3C - Accent / Negative KPI
COLOR_GREEN = RGBColor(39, 174, 96)    # #27AE60 - Positive KPI
COLOR_BLUE = RGBColor(41, 128, 185)    # #2980B9 - Secondary Accent
COLOR_CARD_BG = RGBColor(248, 250, 252)# #F8FAFC - Card Fill
COLOR_CARD_BORDER = RGBColor(226, 232, 240) # #E2E8F0
COLOR_WHITE = RGBColor(255, 255, 255)

def add_header(slide, title_text, category_text="SHOPVISTA COMMERCE • USE CASE 06"):
    """Adds a standard executive header to a slide."""
    header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.9))
    tf = header_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    p_cat = tf.paragraphs[0]
    p_cat.text = category_text.upper()
    p_cat.font.name = "Arial"
    p_cat.font.size = Pt(10)
    p_cat.font.bold = True
    p_cat.font.color.rgb = COLOR_CORAL

    p_title = tf.add_paragraph()
    p_title.text = title_text
    p_title.font.name = "Arial"
    p_title.font.size = Pt(22)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_NAVY
    p_title.space_before = Pt(3)

def add_card(slide, left, top, width, height, bg_color=COLOR_CARD_BG, border_color=COLOR_CARD_BORDER):
    """Draws a rounded rectangular container card."""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
    else:
        card.line.fill.background()
    return card

# ==============================================================================
# SLIDE 1: TITLE SLIDE
# ==============================================================================
slide1 = prs.slides.add_slide(blank_slide_layout)

# Background Accent Banner
bg_banner = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
bg_banner.fill.solid()
bg_banner.fill.fore_color.rgb = COLOR_NAVY
bg_banner.line.fill.background()

# Title text box
tb1 = slide1.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(11.0), Inches(3.8))
tf1 = tb1.text_frame
tf1.word_wrap = True

p_sub0 = tf1.paragraphs[0]
p_sub0.text = "DS SPRINT ASSESSMENT • SKILL BASED ASSESSMENT 2 (U23AD491)"
p_sub0.font.name = "Arial"
p_sub0.font.size = Pt(13)
p_sub0.font.bold = True
p_sub0.font.color.rgb = COLOR_CORAL

p_main = tf1.add_paragraph()
p_main.text = "E-Commerce Return Insight"
p_main.font.name = "Arial"
p_main.font.size = Pt(38)
p_main.font.bold = True
p_main.font.color.rgb = COLOR_WHITE
p_main.space_before = Pt(8)

p_tag = tf1.add_paragraph()
p_tag.text = "Understanding Root Causes, Quantifying Return Elasticity, and Recommending Interventions to Cut Avoidable Returns"
p_tag.font.name = "Arial"
p_tag.font.size = Pt(16)
p_tag.font.color.rgb = RGBColor(203, 213, 225)
p_tag.space_before = Pt(10)

# Metadata card on bottom
tb_meta = slide1.shapes.add_textbox(Inches(1.2), Inches(5.3), Inches(11.0), Inches(1.5))
tf_meta = tb_meta.text_frame
tf_meta.word_wrap = True

p_m1 = tf_meta.paragraphs[0]
p_m1.text = "Organization: ShopVista Commerce Pvt. Ltd. | Use Case 06: E-Commerce Return Insight"
p_m1.font.name = "Arial"
p_m1.font.size = Pt(12)
p_m1.font.bold = True
p_m1.font.color.rgb = COLOR_WHITE

p_m2 = tf_meta.add_paragraph()
p_m2.text = "Dataset Notice: Synthetic Assessment Data (5,000 distinct orders analyzed) | Author: Candidate Data Scientist"
p_m2.font.name = "Arial"
p_m2.font.size = Pt(11)
p_m2.font.color.rgb = RGBColor(148, 163, 184)
p_m2.space_before = Pt(4)

# ==============================================================================
# SLIDE 2: BUSINESS CONTEXT & PROBLEM OBJECTIVE
# ==============================================================================
slide2 = prs.slides.add_slide(blank_slide_layout)
add_header(slide2, "Business Problem, Objectives, and Core Questions")

# 3 Left KPI / Pillar Cards
pillar_data = [
    ("12.78%", "Catalog Return Rate", "639 total product returns out of 5,000 clean orders.", COLOR_NAVY),
    ("₹35.8L", "Gross Return Merchandise Value", "Substantial working capital locked in reverse transit.", COLOR_CORAL),
    ("28.3%", "Primary Reason: Sizing", "181 returns directly caused by size/fit mismatches.", COLOR_BLUE)
]

for idx, (val, title, desc, col) in enumerate(pillar_data):
    top_pos = Inches(1.5 + idx * 1.8)
    add_card(slide2, Inches(0.8), top_pos, Inches(3.6), Inches(1.6))
    tb = slide2.shapes.add_textbox(Inches(1.0), top_pos + Inches(0.15), Inches(3.2), Inches(1.3))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p_v = tf.paragraphs[0]
    p_v.text = val
    p_v.font.name = "Arial"
    p_v.font.size = Pt(26)
    p_v.font.bold = True
    p_v.font.color.rgb = col
    
    p_t = tf.add_paragraph()
    p_t.text = title
    p_t.font.name = "Arial"
    p_t.font.size = Pt(11)
    p_t.font.bold = True
    p_t.font.color.rgb = COLOR_DARK
    p_t.space_before = Pt(2)

    p_d = tf.add_paragraph()
    p_d.text = desc
    p_d.font.name = "Arial"
    p_d.font.size = Pt(10)
    p_d.font.color.rgb = COLOR_MUTED
    p_d.space_before = Pt(2)

# Right Context Card
add_card(slide2, Inches(4.7), Inches(1.5), Inches(7.8), Inches(5.2))
tb_right = slide2.shapes.add_textbox(Inches(5.0), Inches(1.8), Inches(7.2), Inches(4.6))
tf_right = tb_right.text_frame
tf_right.word_wrap = True

p_r0 = tf_right.paragraphs[0]
p_r0.text = "Executive Challenge & Analytical Mandate"
p_r0.font.name = "Arial"
p_r0.font.size = Pt(15)
p_r0.font.bold = True
p_r0.font.color.rgb = COLOR_NAVY

p_r1 = tf_right.add_paragraph()
p_r1.text = "In e-commerce retailing, product returns represent an acute operational drag. Each return incurs two-way shipping fees, warehouse handling, item degradation, and inventory obsolescence. ShopVista Commerce commissioned this investigation to resolve five key questions:"
p_r1.font.name = "Arial"
p_r1.font.size = Pt(11)
p_r1.font.color.rgb = COLOR_DARK
p_r1.space_before = Pt(6)

bullets = [
    ("Category Concentration:", "Which specific merchandise departments exhibit disproportionate return rates?"),
    ("Discount Elasticity:", "Does aggressive price discounting induce return-prone impulse purchasing?"),
    ("Fulfillment Service Levels:", "How severely do late logistics deliveries escalate return propensity?"),
    ("Merchant Disparities:", "Are returns systematically concentrated among specific third-party marketplace sellers?"),
    ("Customer Lifecycle:", "Do repeat, high-frequency patrons demonstrate greater retention compared to new shoppers?")
]

for title_b, text_b in bullets:
    p_b = tf_right.add_paragraph()
    p_b.text = f"•  {title_b} "
    p_b.font.name = "Arial"
    p_b.font.size = Pt(11)
    p_b.font.bold = True
    p_b.font.color.rgb = COLOR_NAVY
    p_b.space_before = Pt(6)

    run = p_b.add_run()
    run.text = text_b
    run.font.bold = False
    run.font.color.rgb = COLOR_DARK

# ==============================================================================
# SLIDE 3: DATA PREPARATION & CLEANING AUDIT
# ==============================================================================
slide3 = prs.slides.add_slide(blank_slide_layout)
add_header(slide3, "Synthetic Dataset Architecture and Data Preparation Audit")

# Table summarizing cleaning
table_shape = slide3.shapes.add_table(6, 5, Inches(0.8), Inches(1.5), Inches(11.733), Inches(3.2))
tbl = table_shape.table
tbl.columns[0].width = Inches(1.8)
tbl.columns[1].width = Inches(1.3)
tbl.columns[2].width = Inches(1.3)
tbl.columns[3].width = Inches(2.8)
tbl.columns[4].width = Inches(4.533)

headers_table = ["File Name", "Raw Rows", "Clean Rows", "Anomalies Detected", "Cleaning & Transformation Rule"]
for i, h in enumerate(headers_table):
    cell = tbl.cell(0, i)
    cell.text = h
    cell.fill.solid()
    cell.fill.fore_color.rgb = COLOR_NAVY
    for p in cell.text_frame.paragraphs:
        p.font.name = "Arial"
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = COLOR_WHITE

table_rows_data = [
    ("orders.csv", "5,050", "5,000", "50 exact dups; 18 extreme discounts", "Deduplicated by order_id. Imputed 18 invalid discounts (<0% or >100%) with median (13.18%)."),
    ("products.csv", "450", "450", "Inconsistent sizing; 10 null brands", "Harmonized casing/labels into standard enum (XS to XXL). Filled missing brands as Generic."),
    ("customers.csv", "1,800", "1,800", "15 null regions", "Imputed missing geographic regions with 'Unknown'. Verified customer history tiers."),
    ("delivery.csv", "5,000", "5,000", "5 unparseable; 7 pre-order delivery dates", "Flagged 12 invalid records as 'Invalid Date'. Calculated delays only on 4,988 valid deliveries."),
    ("returns.csv", "5,000", "5,000", "85 missing return reasons on returns", "Imputed 'Unknown' for 85 missing returned reasons. Left non-returns as 'Not Returned' (0).")
]

for row_idx, r_data in enumerate(table_rows_data, start=1):
    for col_idx, val in enumerate(r_data):
        cell = tbl.cell(row_idx, col_idx)
        cell.text = val
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_CARD_BG if row_idx % 2 == 1 else COLOR_WHITE
        for p in cell.text_frame.paragraphs:
            p.font.name = "Arial"
            p.font.size = Pt(9.5)
            p.font.color.rgb = COLOR_DARK
            if col_idx in [1, 2]:
                p.alignment = PP_ALIGN.RIGHT

# Callout banner at bottom
add_card(slide3, Inches(0.8), Inches(5.1), Inches(11.733), Inches(1.8), bg_color=RGBColor(241, 245, 249))
tb_audit = slide3.shapes.add_textbox(Inches(1.0), Inches(5.2), Inches(11.333), Inches(1.5))
tf_audit = tb_audit.text_frame
tf_audit.word_wrap = True

p_a0 = tf_audit.paragraphs[0]
p_a0.text = "METHODOLOGICAL INTEGRITY & LEAKAGE PREVENTION PRINCIPLES"
p_a0.font.name = "Arial"
p_a0.font.size = Pt(11)
p_a0.font.bold = True
p_a0.font.color.rgb = COLOR_NAVY

audit_notes = (
    "1. Strict Join Grain: Merged 1:1 on order_id. Zero Cartesian duplication. Verified 5,000 distinct final transaction records.\n"
    "2. No Date Fabrication: Invalid delivery records were flagged and isolated rather than assigned arbitrary synthetic dates.\n"
    "3. Target Leakage Prevention: Outcome attributes (return_reason, return_date) were strictly isolated from predictive feature sets."
)
p_a1 = tf_audit.add_paragraph()
p_a1.text = audit_notes
p_a1.font.name = "Arial"
p_a1.font.size = Pt(10)
p_a1.font.color.rgb = COLOR_DARK
p_a1.space_before = Pt(4)

# ==============================================================================
# SLIDE 4: FINDING 1 — CATEGORY RETURN RATES & SIZING MISMATCH
# ==============================================================================
slide4 = prs.slides.add_slide(blank_slide_layout)
add_header(slide4, "Finding 1: Apparel Sizing Mismatch is the Dominant Return Driver")

# Left Image: Chart 1
img_cat_path = os.path.join(FIG_DIR, "eda_1_return_rate_by_category.png")
if os.path.exists(img_cat_path):
    slide4.shapes.add_picture(img_cat_path, Inches(0.8), Inches(1.5), width=Inches(6.2))

# Right Insight Card
add_card(slide4, Inches(7.3), Inches(1.5), Inches(5.2), Inches(5.4))
tb_f1 = slide4.shapes.add_textbox(Inches(7.5), Inches(1.7), Inches(4.8), Inches(5.0))
tf_f1 = tb_f1.text_frame
tf_f1.word_wrap = True

p_f1_0 = tf_f1.paragraphs[0]
p_f1_0.text = "Key Empirical Evidence & Diagnostics"
p_f1_0.font.name = "Arial"
p_f1_0.font.size = Pt(14)
p_f1_0.font.bold = True
p_f1_0.font.color.rgb = COLOR_NAVY

f1_points = [
    ("Apparel Outlier Severity:", "Apparel return rate is 19.53% (198 returns / 1,014 orders), which is 52.8% higher than catalog average (12.78%)."),
    ("Footwear Alignment:", "Footwear exhibits the second highest return rate at 14.15% (90 returns / 636 orders), confirming wearable categories suffer disproportionately."),
    ("Non-Sized Resiliency:", "Electronics (8.60%) and Beauty (8.57%) display return rates under 9%, demonstrating that standardized, specification-driven categories perform stably."),
    ("Root Cause Driver:", "Customer return feedback shows 'Size/Fit Issue' accounts for 181 returns (28.33% of all returns)—the single largest category in the business."),
    ("Business Implication:", "Lack of standardized digital sizing and inconsistent brand tailoring force customers into 'bracketing' (ordering multiple sizes and returning ill-fitting items).")
]

for title_p, text_p in f1_points:
    p_pt = tf_f1.add_paragraph()
    p_pt.text = f"• {title_p} "
    p_pt.font.name = "Arial"
    p_pt.font.size = Pt(10.5)
    p_pt.font.bold = True
    p_pt.font.color.rgb = COLOR_CORAL if "Outlier" in title_p or "Root" in title_p else COLOR_NAVY
    p_pt.space_before = Pt(6)

    run = p_pt.add_run()
    run.text = text_p
    run.font.bold = False
    run.font.color.rgb = COLOR_DARK

# ==============================================================================
# SLIDE 5: FINDING 2 — DELIVERY DELAYS ESCALATE RETURNS
# ==============================================================================
slide5 = prs.slides.add_slide(blank_slide_layout)
add_header(slide5, "Finding 2: Fulfillment Delays Drive a 67.5% Surge in Return Propensity")

# Left Image: Chart 3
img_del_path = os.path.join(FIG_DIR, "eda_3_delivery_delay_vs_return_rate.png")
if os.path.exists(img_del_path):
    slide5.shapes.add_picture(img_del_path, Inches(0.8), Inches(1.5), width=Inches(6.2))

# Right Insight Card
add_card(slide5, Inches(7.3), Inches(1.5), Inches(5.2), Inches(5.4))
tb_f2 = slide5.shapes.add_textbox(Inches(7.5), Inches(1.7), Inches(4.8), Inches(5.0))
tf_f2 = tb_f2.text_frame
tf_f2.word_wrap = True

p_f2_0 = tf_f2.paragraphs[0]
p_f2_0.text = "Fulfillment SLA Impact on Customer Retention"
p_f2_0.font.name = "Arial"
p_f2_0.font.size = Pt(14)
p_f2_0.font.bold = True
p_f2_0.font.color.rgb = COLOR_NAVY

f2_points = [
    ("On-Time vs Late Disparity:", "Deliveries on-time or early return at 11.04% (420 returns / 3,804 orders). Delayed deliveries spike to 18.50% (219 returns / 1,184 orders)—a 67.5% relative increase."),
    ("Nonlinear Delay Compounding:", "While 1-day delays return at 14.83%, extreme delays of 5 days post-promise push returns to 24.34% (nearly 1 in every 4 packages rejected)."),
    ("Second Largest Return Reason:", "'Late Delivery' is cited as the primary reason in 135 returned orders (21.13% of total returns), confirming logistics friction directly provokes cancellation."),
    ("Customer Need-State Decay:", "When merchandise arrives after the expected window, the customer's immediate need or gifting occasion expires, or alternative items are purchased locally."),
    ("Actionable Takeaway:", "Tightening courier dispatch SLA thresholds and improving tracking updates can salvage up to 80+ avoidable returns annually.")
]

for title_p, text_p in f2_points:
    p_pt = tf_f2.add_paragraph()
    p_pt.text = f"• {title_p} "
    p_pt.font.name = "Arial"
    p_pt.font.size = Pt(10.5)
    p_pt.font.bold = True
    p_pt.font.color.rgb = COLOR_CORAL if "Disparity" in title_p or "Compounding" in title_p else COLOR_NAVY
    p_pt.space_before = Pt(6)

    run = p_pt.add_run()
    run.text = text_p
    run.font.bold = False
    run.font.color.rgb = COLOR_DARK

# ==============================================================================
# SLIDE 6: FINDING 3 & 4 — DISCOUNT DEPTH & SELLER DISPARITIES
# ==============================================================================
slide6 = prs.slides.add_slide(blank_slide_layout)
add_header(slide6, "Commercial Dynamics: Discount Elasticity and Seller Performance")

# Left Top: Chart 2
img_disc_path = os.path.join(FIG_DIR, "eda_2_discount_vs_return_rate.png")
if os.path.exists(img_disc_path):
    slide6.shapes.add_picture(img_disc_path, Inches(0.8), Inches(1.5), width=Inches(5.7))

# Right: Chart 4
img_seller_path = os.path.join(FIG_DIR, "eda_4_seller_return_rates.png")
if os.path.exists(img_seller_path):
    slide6.shapes.add_picture(img_seller_path, Inches(6.8), Inches(1.5), width=Inches(5.7))

# Bottom Comparative Card
add_card(slide6, Inches(0.8), Inches(5.1), Inches(11.733), Inches(1.9), bg_color=RGBColor(241, 245, 249))
tb_comb = slide6.shapes.add_textbox(Inches(1.0), Inches(5.2), Inches(11.333), Inches(1.7))
tf_comb = tb_comb.text_frame
tf_comb.word_wrap = True

p_c0 = tf_comb.paragraphs[0]
p_c0.text = "DIAGNOSTIC SYNTHESIS: COMMERCIAL INCENTIVES & MERCHANT QUALITY"
p_c0.font.name = "Arial"
p_c0.font.size = Pt(11)
p_c0.font.bold = True
p_c0.font.color.rgb = COLOR_NAVY

comb_text = (
    "• Discount Propensity Gradient: Return rates steadily escalate with promotional depth—from 10.74% in 0-10% discount orders, "
    "to 12.97% (10-20%), 14.93% (20-30%), and 16.72% in 30-50% discounts. Deep markdowns attract impulse buyers with weaker purchase commitment.\n"
    "• Chronic Seller Outliers: Substantial seller disparity exists. High-volume partners SELLER_015 (25.69% returns across 109 orders) and "
    "SELLER_053 (23.62% returns across 127 orders) return at twice the baseline average, indicating catalog listing misrepresentation or defect patterns."
)
p_c1 = tf_comb.add_paragraph()
p_c1.text = comb_text
p_c1.font.name = "Arial"
p_c1.font.size = Pt(10)
p_c1.font.color.rgb = COLOR_DARK
p_c1.space_before = Pt(4)

# ==============================================================================
# SLIDE 7: MACHINE LEARNING & PREDICTIVE INTELLIGENCE
# ==============================================================================
slide7 = prs.slides.add_slide(blank_slide_layout)
add_header(slide7, "Machine Learning Diagnostic: Balanced Random Forest Model (t=0.45)")

# Left: Confusion Matrix Image
img_cm_path = os.path.join(FIG_DIR, "ml_confusion_matrix.png")
if os.path.exists(img_cm_path):
    slide7.shapes.add_picture(img_cm_path, Inches(0.8), Inches(1.5), width=Inches(4.9))

# Right: Model Details & Verified Metrics
add_card(slide7, Inches(6.0), Inches(1.5), Inches(6.5), Inches(5.5))
tb_ml = slide7.shapes.add_textbox(Inches(6.2), Inches(1.65), Inches(6.1), Inches(5.2))
tf_ml = tb_ml.text_frame
tf_ml.word_wrap = True

p_ml0 = tf_ml.paragraphs[0]
p_ml0.text = "Final Model Configuration & Evaluation Truth (Threshold = 0.45)"
p_ml0.font.name = "Arial"
p_ml0.font.size = Pt(13)
p_ml0.font.bold = True
p_ml0.font.color.rgb = COLOR_NAVY

ml_notes = [
    ("Prediction Time & Leakage Prevention:", "Strictly pre-dispatch. Actual delivery delay, transit days, and post-outcome return fields are completely excluded to eliminate target leakage."),
    ("Final Architecture & Configuration:", "Balanced Random Forest (250 trees, max depth 8, min leaf 4, class_weight='balanced') fitted inside a ColumnTransformer preprocessing pipeline."),
    ("Verified Test Metrics (t=0.45):", "Accuracy: 68.50% | Precision: 17.87% | Recall: 40.62% | F1-Score: 0.2482 | ROC-AUC: 0.5607 | PR-AUC: 0.1572 on untouched holdout test set (N=1,000)."),
    ("Confusion Matrix Truth:", "TN = 633, FP = 239, FN = 76, TP = 52. Successfully captures 52 out of 128 actual returns, compared to 0 for the naive majority baseline (0.0% recall)."),
    ("Honest Data Science Assessment:", "Precision is bounded at 17.87% due to 12.78% base prevalence and the fact that no pre-dispatch cohort exceeds ~26% return rate without post-delivery leakage. Fabricating higher scores would violate data integrity."),
    ("Operational Utility at Scale:", "Provides high practical utility for targeted pre-shipping size confirmations on Apparel orders and selecting priority couriers before warehouse dispatch.")
]

for title_m, text_m in ml_notes:
    p_m = tf_ml.add_paragraph()
    p_m.text = f"• {title_m} "
    p_m.font.name = "Arial"
    p_m.font.size = Pt(9.5)
    p_m.font.bold = True
    p_m.font.color.rgb = COLOR_NAVY
    p_m.space_before = Pt(4)

    run = p_m.add_run()
    run.text = text_m
    run.font.bold = False
    run.font.color.rgb = COLOR_DARK



# ==============================================================================
# SLIDE 8: BUSINESS IMPLICATIONS & MANAGEMENT FOCUS
# ==============================================================================
slide8 = prs.slides.add_slide(blank_slide_layout)
add_header(slide8, "Business Implications & Areas Requiring Management Attention")

cards_impact = [
    ("Reverse Logistics Margin Drag",
     "Across 639 returned orders, reverse freight, restocking, and packaging consume an estimated ₹9.5L–₹12.8L annually in operational overhead, eroding net EBITDA margins.",
     COLOR_CORAL),
    ("Apparel Sizing Dissatisfaction",
     "With 19.53% of Apparel returned, customer frustration with fit consistency reduces long-term customer lifetime value (LTV) and drives churn among first-time shoppers.",
     COLOR_NAVY),
    ("Fulfillment Courier SLA Inefficiencies",
     "1,184 packages arrived past promised dates, triggering 219 returns. Courier partners failing SLA guidelines directly inflate return volumes and customer support tickets.",
     COLOR_BLUE),
    ("Marketplace Seller Quality Governance",
     "A small cohort of marketplace sellers exhibit return rates exceeding 25%. Without quality gating, poor vendor catalogs degrade overall ShopVista platform trust.",
     RGBColor(211, 84, 0))
]

for idx, (title_c, desc_c, color_c) in enumerate(cards_impact):
    x = Inches(0.8 + (idx % 2) * 5.95)
    y = Inches(1.5 + (idx // 2) * 2.7)
    add_card(slide8, x, y, Inches(5.75), Inches(2.5))
    
    tb = slide8.shapes.add_textbox(x + Inches(0.2), y + Inches(0.2), Inches(5.35), Inches(2.1))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p_t = tf.paragraphs[0]
    p_t.text = title_c
    p_t.font.name = "Arial"
    p_t.font.size = Pt(13)
    p_t.font.bold = True
    p_t.font.color.rgb = color_c
    
    p_d = tf.add_paragraph()
    p_d.text = desc_c
    p_d.font.name = "Arial"
    p_d.font.size = Pt(10.5)
    p_d.font.color.rgb = COLOR_DARK
    p_d.space_before = Pt(6)

# ==============================================================================
# SLIDE 9: STRATEGIC RECOMMENDATIONS & ACTION PLAN
# ==============================================================================
slide9 = prs.slides.add_slide(blank_slide_layout)
add_header(slide9, "Practical, Data-Driven Recommendations to Reduce Avoidable Returns")

recs = [
    ("WORKSTREAM 1: PRODUCT & SIZING INTELLIGENCE",
     "Standardize Apparel Fit & Virtual Size Recommenders",
     "• Deploy an interactive size recommendation widget on product pages using customer height/weight.\n"
     "• Mandate rigorous, standardized measurement charts across all vendors (chest, waist, inseam in cm).\n"
     "• Encourage verified customer size reviews ('Fits True to Size' badge) to prevent bracketing.",
     COLOR_NAVY),
    ("WORKSTREAM 2: LOGISTICS & FULFILLMENT RIGOR",
     "Enforce Courier Dispatch SLAs & Dynamic In-Transit Alerts",
     "• Impose contractual SLA penalties on courier partners exceeding promised delivery dates.\n"
     "• Trigger proactive SMS/WhatsApp tracking updates with automated delay compensation credits.\n"
     "• Prioritize expedited fulfillment routing for high-value orders and repeat customers.",
     COLOR_CORAL),
    ("WORKSTREAM 3: COMMERCIAL & SELLER GOVERNANCE",
     "Seller Performance Gating & Discount Return Reserves",
     "• Establish a 20% return rate threshold: place repeat offenders (SELLER_015, SELLER_053) under catalog audit.\n"
     "• Require high-return sellers to finance reverse logistics surcharges for defective/misdescribed goods.\n"
     "• Implement a structured return reserve on clearance discounts (>30%) to safeguard margins.",
     COLOR_BLUE)
]

for idx, (ws_tag, ws_title, ws_bullets, ws_color) in enumerate(recs):
    x = Inches(0.8 + idx * 4.0)
    add_card(slide9, x, Inches(1.5), Inches(3.75), Inches(5.4))
    
    tb = slide9.shapes.add_textbox(x + Inches(0.2), Inches(1.7), Inches(3.35), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p_tag = tf.paragraphs[0]
    p_tag.text = ws_tag
    p_tag.font.name = "Arial"
    p_tag.font.size = Pt(9.5)
    p_tag.font.bold = True
    p_tag.font.color.rgb = ws_color
    
    p_title = tf.add_paragraph()
    p_title.text = ws_title
    p_title.font.name = "Arial"
    p_title.font.size = Pt(13)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_DARK
    p_title.space_before = Pt(4)
    
    p_body = tf.add_paragraph()
    p_body.text = ws_bullets
    p_body.font.name = "Arial"
    p_body.font.size = Pt(10)
    p_body.font.color.rgb = COLOR_DARK
    p_body.space_before = Pt(8)

# ==============================================================================
# SLIDE 10: CONCLUSION & SUMMARY SCORECARD
# ==============================================================================
slide10 = prs.slides.add_slide(blank_slide_layout)
add_header(slide10, "Executive Conclusion & Operational Scorecard")

# Scorecard summary card
add_card(slide10, Inches(0.8), Inches(1.5), Inches(11.733), Inches(4.2), bg_color=COLOR_WHITE)
tb_sc = slide10.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(11.133), Inches(3.6))
tf_sc = tb_sc.text_frame
tf_sc.word_wrap = True

p_sc0 = tf_sc.paragraphs[0]
p_sc0.text = "Summary Assessment & Strategic Roadmap"
p_sc0.font.name = "Arial"
p_sc0.font.size = Pt(15)
p_sc0.font.bold = True
p_sc0.font.color.rgb = COLOR_NAVY

sc_content = (
    "1. Comprehensive Empirical Diagnostics:\n"
    "   Analysis of 5,000 clean transactions confirmed that product returns at ShopVista are highly concentrated in "
    "wearable merchandise (Apparel at 19.53%) and fulfillment-delayed shipments (18.50% return rate for late orders).\n\n"
    "2. High-Impact Interventions:\n"
    "   Targeting sizing accuracy and courier SLA fulfillment addresses the two largest cited return drivers "
    "(Size/Fit: 28.33%, Late Delivery: 21.13%), accounting for nearly half of all returns on the platform.\n\n"
    "3. Commercial Governance:\n"
    "   Vendor accountability (SELLER_015, SELLER_053) and disciplined promotional discount structures "
    "provide immediate opportunities to preserve gross margins without depressing sales volume.\n\n"
    "4. Intelligence Layer Deployment:\n"
    "   A production ML decision pipeline scoring pre-dispatch risk allows proactive intervention (fit verification prompts, "
    "courier prioritization) before expensive fulfillment occurs."
)
p_sc1 = tf_sc.add_paragraph()
p_sc1.text = sc_content
p_sc1.font.name = "Arial"
p_sc1.font.size = Pt(10.5)
p_sc1.font.color.rgb = COLOR_DARK
p_sc1.space_before = Pt(6)

# Bottom Thank You / Sign-off
add_card(slide10, Inches(0.8), Inches(5.9), Inches(11.733), Inches(1.0), bg_color=COLOR_NAVY)
tb_ty = slide10.shapes.add_textbox(Inches(1.0), Inches(6.0), Inches(11.333), Inches(0.8))
tf_ty = tb_ty.text_frame
tf_ty.word_wrap = True
p_ty = tf_ty.paragraphs[0]
p_ty.text = "Thank You. ShopVista Return Insight Diagnostic Complete. Ready for Management Discussion & Rollout."
p_ty.font.name = "Arial"
p_ty.font.size = Pt(12)
p_ty.font.bold = True
p_ty.font.color.rgb = COLOR_WHITE
p_ty.alignment = PP_ALIGN.CENTER

prs.save(OUTPUT_PPTX)
print(f"PowerPoint successfully created at: {OUTPUT_PPTX}")

# Attempt export to PDF via PowerPoint COM automation on Windows
try:
    import win32com.client
    powerpoint = win32com.client.Dispatch("PowerPoint.Application")
    powerpoint.Visible = 1
    deck = powerpoint.Presentations.Open(os.path.abspath(OUTPUT_PPTX))
    deck.SaveAs(os.path.abspath(OUTPUT_PDF), 32) # 32 = ppSaveAsPDF
    deck.Close()
    powerpoint.Quit()
    print(f"PDF successfully exported at: {OUTPUT_PDF}")
except Exception as e:
    print(f"Could not export to PDF via win32com directly ({e}). Will run via PowerShell COM.")
