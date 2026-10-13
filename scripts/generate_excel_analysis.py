"""
generate_excel_analysis.py
Generates the comprehensive Excel analysis workbook: outputs/ecommerce_return_analysis.xlsx
Using openpyxl with corporate styling, KPI cards, Excel formulas, and native charts.
"""

import os
import pandas as pd
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList

# Paths
DATA_PATH = os.path.join("data", "processed", "ecommerce_returns_final_clean.csv")
OUTPUT_PATH = os.path.join("outputs", "ecommerce_return_analysis.xlsx")
os.makedirs("outputs", exist_ok=True)

df = pd.read_csv(DATA_PATH)

wb = openpyxl.Workbook()
# Remove default sheet
wb.remove(wb.active)

# Color Palette & Styles
NAVY_HEADER = "1B365D"
TEAL_ACCENT = "008080"
LIGHT_BG = "F4F6F9"
CARD_BG = "EBF2FA"
BORDER_GRAY = "D1D5DB"
WHITE = "FFFFFF"

font_title = Font(name="Calibri", size=16, bold=True, color="1B365D")
font_subtitle = Font(name="Calibri", size=11, italic=True, color="4B5563")
font_section = Font(name="Calibri", size=13, bold=True, color="1B365D")
font_header = Font(name="Calibri", size=11, bold=True, color=WHITE)
font_bold = Font(name="Calibri", size=11, bold=True, color="111827")
font_regular = Font(name="Calibri", size=11, color="111827")
font_kpi_title = Font(name="Calibri", size=9, bold=True, color="4B5563")
font_kpi_value = Font(name="Calibri", size=18, bold=True, color="1B365D")

fill_header = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
fill_subtotal = PatternFill(start_color="E5E7EB", end_color="E5E7EB", fill_type="solid")
fill_kpi = PatternFill(start_color=CARD_BG, end_color=CARD_BG, fill_type="solid")

thin_border_side = Side(border_style="thin", color=BORDER_GRAY)
thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
thick_bottom = Border(bottom=Side(border_style="medium", color=NAVY_HEADER))
kpi_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

align_center = Alignment(horizontal="center", vertical="center")
align_left = Alignment(horizontal="left", vertical="center")
align_right = Alignment(horizontal="right", vertical="center")

# ==============================================================================
# SHEET 1: SUMMARY DASHBOARD
# ==============================================================================
ws_dash = wb.create_sheet(title="Summary Dashboard")
ws_dash.views.sheetView[0].showGridLines = True

# Title block
ws_dash["B2"] = "ShopVista Commerce — E-Commerce Return Insight Executive Dashboard"
ws_dash["B2"].font = font_title
ws_dash["B3"] = "Enterprise Returns Diagnostic, Root Causes, and Fulfillment Performance Analysis | Use Case 06"
ws_dash["B3"].font = font_subtitle

# KPI Cards row
kpis = [
    ("Total Order Volume", "=COUNT('Cleaned Data'!A2:A5001)", "#,##0", "B5", "B6", "C6"),
    ("Total Product Returns", "=SUM('Cleaned Data'!T2:T5001)", "#,##0", "D5", "D6", "E6"),
    ("Overall Return Rate", "=D6/B6", "0.00%", "F5", "F6", "G6"),
    ("Gross Merchandise Value", "=SUM('Cleaned Data'!W2:W5001)", "₹#,##0.00", "H5", "H6", "I6"),
    ("Net Realized Revenue", "=SUM('Cleaned Data'!Y2:Y5001)", "₹#,##0.00", "J5", "J6", "K6"),
    ("Avg Delivery Delay", "=AVERAGE('Cleaned Data'!Q2:Q5001)", "0.00 \"days\"", "L5", "L6", "M6"),
]

for title, formula, num_fmt, top_cell, val_cell, merge_end in kpis:
    ws_dash[top_cell] = title
    ws_dash[top_cell].font = font_kpi_title
    ws_dash[top_cell].alignment = align_center
    ws_dash[top_cell].fill = fill_kpi

    ws_dash[val_cell] = formula
    ws_dash[val_cell].font = font_kpi_value
    ws_dash[val_cell].alignment = align_center
    ws_dash[val_cell].fill = fill_kpi
    ws_dash[val_cell].number_format = num_fmt

    # Apply borders
    for row in ws_dash[f"{top_cell[0]}5:{merge_end[0]}6"]:
        for c in row:
            c.border = thin_border
            c.fill = fill_kpi

# High-Level Summary Tables on Dashboard
# Table 1: Category Return Summary
ws_dash["B9"] = "1. Category Return Performance Summary"
ws_dash["B9"].font = font_section

cat_headers = ["Category", "Orders", "Returns", "Return Rate (%)", "Net Sales (₹)"]
for col_idx, h in enumerate(cat_headers, start=2):
    cell = ws_dash.cell(row=10, column=col_idx, value=h)
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = align_center

# Fill Category summary using formulas referencing Cleaned Data
cats = sorted(df["category"].unique())
for idx, cat in enumerate(cats, start=11):
    ws_dash.cell(row=idx, column=2, value=cat).font = font_regular
    # Orders formula
    ws_dash.cell(row=idx, column=3, value=f"=COUNTIF('Cleaned Data'!$H$2:$H$5001, B{idx})").number_format = "#,##0"
    # Returns formula
    ws_dash.cell(row=idx, column=4, value=f"=COUNTIFS('Cleaned Data'!$H$2:$H$5001, B{idx}, 'Cleaned Data'!$T$2:$T$5001, 1)").number_format = "#,##0"
    # Return rate formula
    ws_dash.cell(row=idx, column=5, value=f"=D{idx}/C{idx}").number_format = "0.00%"
    # Net sales formula
    ws_dash.cell(row=idx, column=6, value=f"=SUMIFS('Cleaned Data'!$Y$2:$Y$5001, 'Cleaned Data'!$H$2:$H$5001, B{idx})").number_format = "₹#,##0.00"

    for c in range(2, 7):
        ws_dash.cell(row=idx, column=c).border = thin_border
        if c in [3, 4, 5, 6]:
            ws_dash.cell(row=idx, column=c).alignment = align_right
        else:
            ws_dash.cell(row=idx, column=c).alignment = align_left

# Total row
tot_row = 11 + len(cats)
ws_dash.cell(row=tot_row, column=2, value="Total / Catalog Avg").font = font_bold
ws_dash.cell(row=tot_row, column=3, value=f"=SUM(C11:C{tot_row-1})").number_format = "#,##0"
ws_dash.cell(row=tot_row, column=4, value=f"=SUM(D11:D{tot_row-1})").number_format = "#,##0"
ws_dash.cell(row=tot_row, column=5, value=f"=D{tot_row}/C{tot_row}").number_format = "0.00%"
ws_dash.cell(row=tot_row, column=6, value=f"=SUM(F11:F{tot_row-1})").number_format = "₹#,##0.00"
for c in range(2, 7):
    ws_dash.cell(row=tot_row, column=c).fill = fill_subtotal
    ws_dash.cell(row=tot_row, column=c).font = font_bold
    ws_dash.cell(row=tot_row, column=c).border = thin_border
    if c >= 3:
        ws_dash.cell(row=tot_row, column=c).alignment = align_right

# Table 2: Key Return Reasons
ws_dash["H9"] = "2. Top Return Reasons (Customer Stated)"
ws_dash["H9"].font = font_section

reason_headers = ["Return Reason", "Return Count", "% of Returns"]
for col_idx, h in enumerate(reason_headers, start=8):
    cell = ws_dash.cell(row=10, column=col_idx, value=h)
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = align_center

reasons_order = [
    "Size/Fit Issue", "Late Delivery", "Unknown", "Not as Described",
    "Damaged Product", "Quality Issue", "Wrong Item Delivered", "Changed Mind"
]
for idx, r in enumerate(reasons_order, start=11):
    ws_dash.cell(row=idx, column=8, value=r).font = font_regular
    ws_dash.cell(row=idx, column=9, value=f"=COUNTIF('Cleaned Data'!$V$2:$V$5001, H{idx})").number_format = "#,##0"
    ws_dash.cell(row=idx, column=10, value=f"=I{idx}/$D$6").number_format = "0.00%"

    for c in range(8, 11):
        ws_dash.cell(row=idx, column=c).border = thin_border
        if c >= 9:
            ws_dash.cell(row=idx, column=c).alignment = align_right
        else:
            ws_dash.cell(row=idx, column=c).alignment = align_left

tot_r_row = 11 + len(reasons_order)
ws_dash.cell(row=tot_r_row, column=8, value="Total Stated Returns").font = font_bold
ws_dash.cell(row=tot_r_row, column=9, value=f"=SUM(I11:I{tot_r_row-1})").number_format = "#,##0"
ws_dash.cell(row=tot_r_row, column=10, value=f"=SUM(J11:J{tot_r_row-1})").number_format = "0.00%"
for c in range(8, 11):
    ws_dash.cell(row=tot_r_row, column=c).fill = fill_subtotal
    ws_dash.cell(row=tot_r_row, column=c).font = font_bold
    ws_dash.cell(row=tot_r_row, column=c).border = thin_border
    if c >= 9:
        ws_dash.cell(row=tot_r_row, column=c).alignment = align_right

# Embed Dashboard Native Chart: Category Return Rate Chart
chart_cat = BarChart()
chart_cat.type = "col"
chart_cat.style = 10
chart_cat.title = "Return Rate by Product Category (%)"
chart_cat.y_axis.title = "Return Rate"
chart_cat.x_axis.title = "Category"
chart_cat.height = 10
chart_cat.width = 16

data_cat = Reference(ws_dash, min_col=5, min_row=10, max_row=tot_row-1)
cats_ref = Reference(ws_dash, min_col=2, min_row=11, max_row=tot_row-1)
chart_cat.add_data(data_cat, titles_from_data=True)
chart_cat.set_categories(cats_ref)
chart_cat.legend = None
ws_dash.add_chart(chart_cat, "B21")

# Embed Dashboard Native Chart: Top Return Reasons
chart_reason = BarChart()
chart_reason.type = "bar"
chart_reason.style = 11
chart_reason.title = "Top Stated Return Reasons (Count)"
chart_reason.x_axis.title = "Returns Count"
chart_reason.y_axis.title = "Reason"
chart_reason.height = 10
chart_reason.width = 16

data_reason = Reference(ws_dash, min_col=9, min_row=10, max_row=tot_r_row-1)
reasons_ref = Reference(ws_dash, min_col=8, min_row=11, max_row=tot_r_row-1)
chart_reason.add_data(data_reason, titles_from_data=True)
chart_reason.set_categories(reasons_ref)
chart_reason.legend = None
ws_dash.add_chart(chart_reason, "H21")

# ==============================================================================
# SHEET 2: CLEANED DATA
# ==============================================================================
ws_clean = wb.create_sheet(title="Cleaned Data")
ws_clean.views.sheetView[0].showGridLines = True

headers = [
    "order_id", "customer_id", "product_id", "order_date", "quantity",
    "price", "discount_pct", "category", "brand", "size",
    "seller", "list_price", "region", "order_history", "promised_date",
    "delivered_date", "delivery_delay_days", "delivery_status", "courier", "return_flag",
    "return_reason", "return_reason_descriptive", "gross_order_value", "discount_amount",
    "net_order_value", "discount_bracket", "transit_days", "is_late_delivery"
]

for col_idx, h in enumerate(headers, start=1):
    cell = ws_clean.cell(row=1, column=col_idx, value=h)
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = align_center

# Populate rows
clean_records = df[headers].copy()
for row_idx, row in enumerate(clean_records.itertuples(index=False), start=2):
    for col_idx, val in enumerate(row, start=1):
        cell = ws_clean.cell(row=row_idx, column=col_idx)
        if pd.isna(val):
            cell.value = ""
        else:
            cell.value = val

        # Number formats
        col_name = headers[col_idx-1]
        if col_name in ["price", "list_price", "gross_order_value", "discount_amount", "net_order_value"]:
            cell.number_format = "₹#,##0.00"
            cell.alignment = align_right
        elif col_name in ["discount_pct"]:
            cell.number_format = "0.00"
            cell.alignment = align_right
        elif col_name in ["quantity", "delivery_delay_days", "transit_days", "is_late_delivery", "return_flag"]:
            if cell.value != "":
                cell.number_format = "#,##0"
            cell.alignment = align_right
        else:
            cell.alignment = align_left

# ==============================================================================
# SHEET 3: RETURN RATE BY CATEGORY
# ==============================================================================
ws_cat = wb.create_sheet(title="Return Rate by Category")
ws_cat.views.sheetView[0].showGridLines = True

ws_cat["B2"] = "Product Return Rate Analysis by Merchandise Category"
ws_cat["B2"].font = font_title
ws_cat["B3"] = "Calculated across 5,000 orders to diagnose department-level return severity"
ws_cat["B3"].font = font_subtitle

cat_table_headers = ["Category", "Total Orders", "Returned Orders", "Return Rate (%)", "Gross Value (₹)", "Discount Value (₹)", "Net Value (₹)"]
for col_idx, h in enumerate(cat_table_headers, start=2):
    cell = ws_cat.cell(row=5, column=col_idx, value=h)
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = align_center

for idx, cat in enumerate(cats, start=6):
    ws_cat.cell(row=idx, column=2, value=cat).font = font_regular
    ws_cat.cell(row=idx, column=3, value=f"=COUNTIF('Cleaned Data'!$H$2:$H$5001, B{idx})").number_format = "#,##0"
    ws_cat.cell(row=idx, column=4, value=f"=COUNTIFS('Cleaned Data'!$H$2:$H$5001, B{idx}, 'Cleaned Data'!$T$2:$T$5001, 1)").number_format = "#,##0"
    ws_cat.cell(row=idx, column=5, value=f"=D{idx}/C{idx}").number_format = "0.00%"
    ws_cat.cell(row=idx, column=6, value=f"=SUMIFS('Cleaned Data'!$W$2:$W$5001, 'Cleaned Data'!$H$2:$H$5001, B{idx})").number_format = "₹#,##0.00"
    ws_cat.cell(row=idx, column=7, value=f"=SUMIFS('Cleaned Data'!$X$2:$X$5001, 'Cleaned Data'!$H$2:$H$5001, B{idx})").number_format = "₹#,##0.00"
    ws_cat.cell(row=idx, column=8, value=f"=SUMIFS('Cleaned Data'!$Y$2:$Y$5001, 'Cleaned Data'!$H$2:$H$5001, B{idx})").number_format = "₹#,##0.00"

    for c in range(2, 9):
        ws_cat.cell(row=idx, column=c).border = thin_border
        if c >= 3:
            ws_cat.cell(row=idx, column=c).alignment = align_right
        else:
            ws_cat.cell(row=idx, column=c).alignment = align_left

tot_c_row = 6 + len(cats)
ws_cat.cell(row=tot_c_row, column=2, value="Total / Catalog Summary").font = font_bold
ws_cat.cell(row=tot_c_row, column=3, value=f"=SUM(C6:C{tot_c_row-1})").number_format = "#,##0"
ws_cat.cell(row=tot_c_row, column=4, value=f"=SUM(D6:D{tot_c_row-1})").number_format = "#,##0"
ws_cat.cell(row=tot_c_row, column=5, value=f"=D{tot_c_row}/C{tot_c_row}").number_format = "0.00%"
ws_cat.cell(row=tot_c_row, column=6, value=f"=SUM(F6:F{tot_c_row-1})").number_format = "₹#,##0.00"
ws_cat.cell(row=tot_c_row, column=7, value=f"=SUM(G6:G{tot_c_row-1})").number_format = "₹#,##0.00"
ws_cat.cell(row=tot_c_row, column=8, value=f"=SUM(H6:H{tot_c_row-1})").number_format = "₹#,##0.00"

for c in range(2, 9):
    ws_cat.cell(row=tot_c_row, column=c).fill = fill_subtotal
    ws_cat.cell(row=tot_c_row, column=c).font = font_bold
    ws_cat.cell(row=tot_c_row, column=c).border = thin_border
    if c >= 3:
        ws_cat.cell(row=tot_c_row, column=c).alignment = align_right

# Chart
cat_bar = BarChart()
cat_bar.type = "col"
cat_bar.style = 10
cat_bar.title = "Category Return Rates"
cat_bar.y_axis.title = "Return Rate"
cat_bar.x_axis.title = "Category"
cat_bar.height = 12
cat_bar.width = 18

data_ref = Reference(ws_cat, min_col=5, min_row=5, max_row=tot_c_row-1)
cats_labels = Reference(ws_cat, min_col=2, min_row=6, max_row=tot_c_row-1)
cat_bar.add_data(data_ref, titles_from_data=True)
cat_bar.set_categories(cats_labels)
cat_bar.legend = None
ws_cat.add_chart(cat_bar, "B16")

# ==============================================================================
# SHEET 4: DISCOUNT COMPARISON
# ==============================================================================
ws_disc = wb.create_sheet(title="Discount Comparison")
ws_disc.views.sheetView[0].showGridLines = True

ws_disc["B2"] = "Discount Tier vs Return Propensity Analysis"
ws_disc["B2"].font = font_title
ws_disc["B3"] = "Evaluating whether promotional discounting depth drives customer returns"
ws_disc["B3"].font = font_subtitle

disc_headers = ["Discount Bracket", "Orders Count", "Returns Count", "Return Rate (%)", "Avg Gross Order (₹)", "Total Discount (₹)"]
for col_idx, h in enumerate(disc_headers, start=2):
    cell = ws_disc.cell(row=5, column=col_idx, value=h)
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = align_center

disc_brackets = ["0-10%", "10-20%", "20-30%", "30-50%", "50%+"]
for idx, brk in enumerate(disc_brackets, start=6):
    ws_disc.cell(row=idx, column=2, value=brk).font = font_regular
    ws_disc.cell(row=idx, column=3, value=f"=COUNTIF('Cleaned Data'!$Z$2:$Z$5001, B{idx})").number_format = "#,##0"
    ws_disc.cell(row=idx, column=4, value=f"=COUNTIFS('Cleaned Data'!$Z$2:$Z$5001, B{idx}, 'Cleaned Data'!$T$2:$T$5001, 1)").number_format = "#,##0"
    ws_disc.cell(row=idx, column=5, value=f"=D{idx}/C{idx}").number_format = "0.00%"
    ws_disc.cell(row=idx, column=6, value=f"=AVERAGEIF('Cleaned Data'!$Z$2:$Z$5001, B{idx}, 'Cleaned Data'!$W$2:$W$5001)").number_format = "₹#,##0.00"
    ws_disc.cell(row=idx, column=7, value=f"=SUMIFS('Cleaned Data'!$X$2:$X$5001, 'Cleaned Data'!$Z$2:$Z$5001, B{idx})").number_format = "₹#,##0.00"

    for c in range(2, 8):
        ws_disc.cell(row=idx, column=c).border = thin_border
        if c >= 3:
            ws_disc.cell(row=idx, column=c).alignment = align_right
        else:
            ws_disc.cell(row=idx, column=c).alignment = align_left

tot_d_row = 6 + len(disc_brackets)
ws_disc.cell(row=tot_d_row, column=2, value="Total / Overall").font = font_bold
ws_disc.cell(row=tot_d_row, column=3, value=f"=SUM(C6:C{tot_d_row-1})").number_format = "#,##0"
ws_disc.cell(row=tot_d_row, column=4, value=f"=SUM(D6:D{tot_d_row-1})").number_format = "#,##0"
ws_disc.cell(row=tot_d_row, column=5, value=f"=D{tot_d_row}/C{tot_d_row}").number_format = "0.00%"
ws_disc.cell(row=tot_d_row, column=6, value=f"=AVERAGE('Cleaned Data'!$W$2:$W$5001)").number_format = "₹#,##0.00"
ws_disc.cell(row=tot_d_row, column=7, value=f"=SUM(G6:G{tot_d_row-1})").number_format = "₹#,##0.00"

for c in range(2, 8):
    ws_disc.cell(row=tot_d_row, column=c).fill = fill_subtotal
    ws_disc.cell(row=tot_d_row, column=c).font = font_bold
    ws_disc.cell(row=tot_d_row, column=c).border = thin_border
    if c >= 3:
        ws_disc.cell(row=tot_d_row, column=c).alignment = align_right

# Chart
disc_chart = LineChart()
disc_chart.title = "Return Rate Across Discount Brackets (%)"
disc_chart.style = 13
disc_chart.y_axis.title = "Return Rate"
disc_chart.x_axis.title = "Discount Bracket"
disc_chart.height = 11
disc_chart.width = 17

data_disc = Reference(ws_disc, min_col=5, min_row=5, max_row=tot_d_row-1)
labels_disc = Reference(ws_disc, min_col=2, min_row=6, max_row=tot_d_row-1)
disc_chart.add_data(data_disc, titles_from_data=True)
disc_chart.set_categories(labels_disc)
disc_chart.legend = None
ws_disc.add_chart(disc_chart, "B14")

# ==============================================================================
# SHEET 5: DELIVERY DELAY ANALYSIS
# ==============================================================================
ws_deliv = wb.create_sheet(title="Delivery Delay Analysis")
ws_deliv.views.sheetView[0].showGridLines = True

ws_deliv["B2"] = "Delivery Delay vs Return Likelihood Diagnostic"
ws_deliv["B2"].font = font_title
ws_deliv["B3"] = "Measuring fulfillment performance against customer return decisions (4,988 valid deliveries)"
ws_deliv["B3"].font = font_subtitle

deliv_headers = ["Delay in Days", "Fulfillment Status", "Order Volume", "Returns", "Return Rate (%)"]
for col_idx, h in enumerate(deliv_headers, start=2):
    cell = ws_deliv.cell(row=5, column=col_idx, value=h)
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = align_center

delay_days_list = [-2, -1, 0, 1, 2, 3, 4, 5]
for idx, d in enumerate(delay_days_list, start=6):
    status_str = "Early (-2d)" if d == -2 else ("Early (-1d)" if d == -1 else ("On-Time (0d)" if d == 0 else f"Late (+{d}d)"))
    ws_deliv.cell(row=idx, column=2, value=d).number_format = "#,##0"
    ws_deliv.cell(row=idx, column=3, value=status_str).font = font_regular
    ws_deliv.cell(row=idx, column=4, value=f"=COUNTIF('Cleaned Data'!$Q$2:$Q$5001, B{idx})").number_format = "#,##0"
    ws_deliv.cell(row=idx, column=5, value=f"=COUNTIFS('Cleaned Data'!$Q$2:$Q$5001, B{idx}, 'Cleaned Data'!$T$2:$T$5001, 1)").number_format = "#,##0"
    ws_deliv.cell(row=idx, column=6, value=f"=E{idx}/D{idx}").number_format = "0.00%"

    for c in range(2, 7):
        ws_deliv.cell(row=idx, column=c).border = thin_border
        if c in [2, 4, 5, 6]:
            ws_deliv.cell(row=idx, column=c).alignment = align_right
        else:
            ws_deliv.cell(row=idx, column=c).alignment = align_left

tot_del_row = 6 + len(delay_days_list)
ws_deliv.cell(row=tot_del_row, column=2, value="Total Valid Deliveries").font = font_bold
ws_deliv.cell(row=tot_del_row, column=3, value="").fill = fill_subtotal
ws_deliv.cell(row=tot_del_row, column=4, value=f"=SUM(D6:D{tot_del_row-1})").number_format = "#,##0"
ws_deliv.cell(row=tot_del_row, column=5, value=f"=SUM(E6:E{tot_del_row-1})").number_format = "#,##0"
ws_deliv.cell(row=tot_del_row, column=6, value=f"=E{tot_del_row}/D{tot_del_row}").number_format = "0.00%"

for c in range(2, 7):
    ws_deliv.cell(row=tot_del_row, column=c).fill = fill_subtotal
    ws_deliv.cell(row=tot_del_row, column=c).font = font_bold
    ws_deliv.cell(row=tot_del_row, column=c).border = thin_border
    if c >= 4 or c == 2:
        ws_deliv.cell(row=tot_del_row, column=c).alignment = align_right

# Delivery Chart
deliv_chart = BarChart()
deliv_chart.type = "col"
deliv_chart.style = 11
deliv_chart.title = "Return Rate by Delivery Delay (Days)"
deliv_chart.y_axis.title = "Return Rate"
deliv_chart.x_axis.title = "Fulfillment Status"
deliv_chart.height = 11
deliv_chart.width = 17

data_deliv = Reference(ws_deliv, min_col=6, min_row=5, max_row=tot_del_row-1)
labels_deliv = Reference(ws_deliv, min_col=3, min_row=6, max_row=tot_del_row-1)
deliv_chart.add_data(data_deliv, titles_from_data=True)
deliv_chart.set_categories(labels_deliv)
deliv_chart.legend = None
ws_deliv.add_chart(deliv_chart, "B17")

# ==============================================================================
# SHEET 6: SELLER COMPARISON
# ==============================================================================
ws_seller = wb.create_sheet(title="Seller Comparison")
ws_seller.views.sheetView[0].showGridLines = True

ws_seller["B2"] = "Seller Partner Return Benchmarking"
ws_seller["B2"].font = font_title
ws_seller["B3"] = "Benchmarking third-party marketplace sellers with significant transaction volume (>= 30 orders)"
ws_seller["B3"].font = font_subtitle

seller_headers = ["Seller ID", "Total Orders", "Returned Orders", "Return Rate (%)", "Net Volume Sold (₹)"]
for col_idx, h in enumerate(seller_headers, start=2):
    cell = ws_seller.cell(row=5, column=col_idx, value=h)
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = align_center

# Top sellers with >= 30 orders
seller_counts = df.groupby("seller")["return_flag"].agg(["count", "sum"])
seller_counts["rate"] = seller_counts["sum"] / seller_counts["count"]
top_sellers_list = seller_counts[seller_counts["count"] >= 30].sort_values(by="rate", ascending=False).head(15).index.tolist()

for idx, s in enumerate(top_sellers_list, start=6):
    ws_seller.cell(row=idx, column=2, value=s).font = font_regular
    ws_seller.cell(row=idx, column=3, value=f"=COUNTIF('Cleaned Data'!$K$2:$K$5001, B{idx})").number_format = "#,##0"
    ws_seller.cell(row=idx, column=4, value=f"=COUNTIFS('Cleaned Data'!$K$2:$K$5001, B{idx}, 'Cleaned Data'!$T$2:$T$5001, 1)").number_format = "#,##0"
    ws_seller.cell(row=idx, column=5, value=f"=D{idx}/C{idx}").number_format = "0.00%"
    ws_seller.cell(row=idx, column=6, value=f"=SUMIFS('Cleaned Data'!$Y$2:$Y$5001, 'Cleaned Data'!$K$2:$K$5001, B{idx})").number_format = "₹#,##0.00"

    for c in range(2, 7):
        ws_seller.cell(row=idx, column=c).border = thin_border
        if c >= 3:
            ws_seller.cell(row=idx, column=c).alignment = align_right
        else:
            ws_seller.cell(row=idx, column=c).alignment = align_left

# Seller Chart
seller_chart = BarChart()
seller_chart.type = "bar"
seller_chart.style = 13
seller_chart.title = "Top Return Sellers (>= 30 Orders)"
seller_chart.x_axis.title = "Return Rate"
seller_chart.y_axis.title = "Seller"
seller_chart.height = 12
seller_chart.width = 17

data_seller = Reference(ws_seller, min_col=5, min_row=5, max_row=5+len(top_sellers_list))
labels_seller = Reference(ws_seller, min_col=2, min_row=6, max_row=5+len(top_sellers_list))
seller_chart.add_data(data_seller, titles_from_data=True)
seller_chart.set_categories(labels_seller)
seller_chart.legend = None
ws_seller.add_chart(seller_chart, "B23")

# ==============================================================================
# SHEET 7: RETURN REASONS
# ==============================================================================
ws_reasons = wb.create_sheet(title="Return Reasons")
ws_reasons.views.sheetView[0].showGridLines = True

ws_reasons["B2"] = "Customer Stated Return Reasons Analysis"
ws_reasons["B2"].font = font_title
ws_reasons["B3"] = "Pareto distribution of reported reasons for all 639 return transactions"
ws_reasons["B3"].font = font_subtitle

rr_headers = ["Return Reason", "Return Count", "% of Total Returns", "Cumulative %"]
for col_idx, h in enumerate(rr_headers, start=2):
    cell = ws_reasons.cell(row=5, column=col_idx, value=h)
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = align_center

for idx, r in enumerate(reasons_order, start=6):
    ws_reasons.cell(row=idx, column=2, value=r).font = font_regular
    ws_reasons.cell(row=idx, column=3, value=f"=COUNTIF('Cleaned Data'!$V$2:$V$5001, B{idx})").number_format = "#,##0"
    ws_reasons.cell(row=idx, column=4, value=f"=C{idx}/$C${6+len(reasons_order)}").number_format = "0.00%"
    if idx == 6:
        ws_reasons.cell(row=idx, column=5, value=f"=D6").number_format = "0.00%"
    else:
        ws_reasons.cell(row=idx, column=5, value=f"=E{idx-1}+D{idx}").number_format = "0.00%"

    for c in range(2, 6):
        ws_reasons.cell(row=idx, column=c).border = thin_border
        if c >= 3:
            ws_reasons.cell(row=idx, column=c).alignment = align_right
        else:
            ws_reasons.cell(row=idx, column=c).alignment = align_left

tot_rr_row = 6 + len(reasons_order)
ws_reasons.cell(row=tot_rr_row, column=2, value="Total Returns Analyzed").font = font_bold
ws_reasons.cell(row=tot_rr_row, column=3, value=f"=SUM(C6:C{tot_rr_row-1})").number_format = "#,##0"
ws_reasons.cell(row=tot_rr_row, column=4, value=f"=SUM(D6:D{tot_rr_row-1})").number_format = "0.00%"
ws_reasons.cell(row=tot_rr_row, column=5, value="100.00%").number_format = "0.00%"

for c in range(2, 6):
    ws_reasons.cell(row=tot_rr_row, column=c).fill = fill_subtotal
    ws_reasons.cell(row=tot_rr_row, column=c).font = font_bold
    ws_reasons.cell(row=tot_rr_row, column=c).border = thin_border
    if c >= 3:
        ws_reasons.cell(row=tot_rr_row, column=c).alignment = align_right

# Reasons Chart
reasons_chart = BarChart()
reasons_chart.type = "col"
reasons_chart.style = 10
reasons_chart.title = "Pareto Breakdown of Return Reasons"
reasons_chart.y_axis.title = "Returns Count"
reasons_chart.x_axis.title = "Return Reason"
reasons_chart.height = 11
reasons_chart.width = 17

data_rr = Reference(ws_reasons, min_col=3, min_row=5, max_row=tot_rr_row-1)
labels_rr = Reference(ws_reasons, min_col=2, min_row=6, max_row=tot_rr_row-1)
reasons_chart.add_data(data_rr, titles_from_data=True)
reasons_chart.set_categories(labels_rr)
reasons_chart.legend = None
ws_reasons.add_chart(reasons_chart, "B16")

# ==============================================================================
# SHEET 8: CUSTOMER SEGMENT ANALYSIS
# ==============================================================================
ws_cust = wb.create_sheet(title="Customer Segment Analysis")
ws_cust.views.sheetView[0].showGridLines = True

ws_cust["B2"] = "Customer Order-History Segment Return Analysis"
ws_cust["B2"].font = font_title
ws_cust["B3"] = "Examining behavioral differences between new, occasional, repeat, and frequent customers"
ws_cust["B3"].font = font_subtitle

cust_headers = ["Customer Segment", "Total Orders", "Returns", "Return Rate (%)", "Avg Order Value (₹)", "Total Net Value (₹)"]
for col_idx, h in enumerate(cust_headers, start=2):
    cell = ws_cust.cell(row=5, column=col_idx, value=h)
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = align_center

segments = ["New", "Occasional", "Repeat", "Frequent"]
for idx, seg in enumerate(segments, start=6):
    ws_cust.cell(row=idx, column=2, value=seg).font = font_regular
    ws_cust.cell(row=idx, column=3, value=f"=COUNTIF('Cleaned Data'!$N$2:$N$5001, B{idx})").number_format = "#,##0"
    ws_cust.cell(row=idx, column=4, value=f"=COUNTIFS('Cleaned Data'!$N$2:$N$5001, B{idx}, 'Cleaned Data'!$T$2:$T$5001, 1)").number_format = "#,##0"
    ws_cust.cell(row=idx, column=5, value=f"=D{idx}/C{idx}").number_format = "0.00%"
    ws_cust.cell(row=idx, column=6, value=f"=AVERAGEIF('Cleaned Data'!$N$2:$N$5001, B{idx}, 'Cleaned Data'!$Y$2:$Y$5001)").number_format = "₹#,##0.00"
    ws_cust.cell(row=idx, column=7, value=f"=SUMIFS('Cleaned Data'!$Y$2:$Y$5001, 'Cleaned Data'!$N$2:$N$5001, B{idx})").number_format = "₹#,##0.00"

    for c in range(2, 8):
        ws_cust.cell(row=idx, column=c).border = thin_border
        if c >= 3:
            ws_cust.cell(row=idx, column=c).alignment = align_right
        else:
            ws_cust.cell(row=idx, column=c).alignment = align_left

tot_cust_row = 6 + len(segments)
ws_cust.cell(row=tot_cust_row, column=2, value="Total / Catalog Summary").font = font_bold
ws_cust.cell(row=tot_cust_row, column=3, value=f"=SUM(C6:C{tot_cust_row-1})").number_format = "#,##0"
ws_cust.cell(row=tot_cust_row, column=4, value=f"=SUM(D6:D{tot_cust_row-1})").number_format = "#,##0"
ws_cust.cell(row=tot_cust_row, column=5, value=f"=D{tot_cust_row}/C{tot_cust_row}").number_format = "0.00%"
ws_cust.cell(row=tot_cust_row, column=6, value=f"=AVERAGE('Cleaned Data'!$Y$2:$Y$5001)").number_format = "₹#,##0.00"
ws_cust.cell(row=tot_cust_row, column=7, value=f"=SUM(G6:G{tot_cust_row-1})").number_format = "₹#,##0.00"

for c in range(2, 8):
    ws_cust.cell(row=tot_cust_row, column=c).fill = fill_subtotal
    ws_cust.cell(row=tot_cust_row, column=c).font = font_bold
    ws_cust.cell(row=tot_cust_row, column=c).border = thin_border
    if c >= 3:
        ws_cust.cell(row=tot_cust_row, column=c).alignment = align_right

# Segment Chart
cust_chart = BarChart()
cust_chart.type = "col"
cust_chart.style = 10
cust_chart.title = "Return Rate by Customer History Segment"
cust_chart.y_axis.title = "Return Rate"
cust_chart.x_axis.title = "Segment"
cust_chart.height = 11
cust_chart.width = 16

data_cust = Reference(ws_cust, min_col=5, min_row=5, max_row=tot_cust_row-1)
labels_cust = Reference(ws_cust, min_col=2, min_row=6, max_row=tot_cust_row-1)
cust_chart.add_data(data_cust, titles_from_data=True)
cust_chart.set_categories(labels_cust)
cust_chart.legend = None
ws_cust.add_chart(cust_chart, "B13")

# Auto-adjust column widths across all sheets
for ws in wb.worksheets:
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len and not val_str.startswith("="):
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

# Specific custom widths for Cleaned Data
ws_clean.column_dimensions["A"].width = 14
ws_clean.column_dimensions["D"].width = 13
ws_clean.column_dimensions["O"].width = 13
ws_clean.column_dimensions["P"].width = 13
ws_clean.column_dimensions["V"].width = 24

wb.save(OUTPUT_PATH)
print(f"Excel workbook successfully created at: {OUTPUT_PATH}")
